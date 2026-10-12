import io

import fitz
import pytest
from docx import Document

from src.parser import MAX_FILE_BYTES, DocumentError, parse_document, parse_path, validate_upload

SAMPLE = "Python developer with PyTorch and scikit-learn experience building NLP models for text classification."


def make_docx_bytes(text: str = SAMPLE) -> bytes:
    document = Document()
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "Docker, Git"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pdf_bytes(text: str = SAMPLE, pages: int = 1) -> bytes:
    pdf = fitz.open()
    for _ in range(pages):
        page = pdf.new_page()
        page.insert_text((72, 72), text)
    data = pdf.tobytes()
    pdf.close()
    return data


def test_parse_docx_extracts_paragraphs_and_tables():
    result = parse_document("cv.docx", make_docx_bytes())
    assert result.ok
    assert "PyTorch" in result.text
    assert "Docker, Git" in result.text


def test_parse_pdf_extracts_text():
    result = parse_document("cv.pdf", make_pdf_bytes())
    assert result.ok
    assert result.num_pages == 1
    assert "scikit-learn" in result.text


def test_unsupported_extension():
    result = parse_document("cv.txt", b"hello world" * 10)
    assert not result.ok
    assert "Unsupported" in result.error


def test_empty_file():
    result = parse_document("cv.pdf", b"")
    assert not result.ok
    assert "empty" in result.error.lower()


def test_corrupt_pdf_with_valid_header():
    result = parse_document("cv.pdf", b"%PDF-1.4\n" + b"garbage" * 50)
    assert not result.ok  # must not raise


def test_fake_pdf_without_header():
    result = parse_document("cv.pdf", b"this is not a pdf" * 10)
    assert not result.ok
    assert "valid PDF" in result.error


def test_fake_docx_not_zip():
    result = parse_document("cv.docx", b"this is not a docx" * 10)
    assert not result.ok
    assert "valid Word" in result.error


def test_docx_zip_but_not_word():
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "hi")
    result = parse_document("cv.docx", buffer.getvalue())
    assert not result.ok


def test_blank_pdf_reports_no_text():
    pdf = fitz.open()
    pdf.new_page()
    data = pdf.tobytes()
    pdf.close()
    result = parse_document("blank.pdf", data)
    assert not result.ok
    assert "No readable text" in result.error


def test_oversized_file_rejected():
    with pytest.raises(DocumentError):
        validate_upload("big.pdf", b"%PDF-" + b"0" * (MAX_FILE_BYTES + 1))


def test_password_protected_pdf():
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), SAMPLE)
    data = pdf.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="secret", owner_pw="owner")
    pdf.close()
    result = parse_document("locked.pdf", data)
    assert not result.ok
    assert "password" in result.error.lower()


def test_parse_path_missing_file(tmp_path):
    result = parse_path(tmp_path / "nope.pdf")
    assert not result.ok
    assert "Could not read" in result.error


def test_parse_path_roundtrip(tmp_path):
    path = tmp_path / "cv.docx"
    path.write_bytes(make_docx_bytes())
    assert parse_path(path).ok
