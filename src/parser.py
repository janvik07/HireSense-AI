"""Document parsing: validate uploads and extract text from PDF / DOCX files.

Design rules:
* ``parse_document`` NEVER raises - every failure becomes ``ParsedDocument.error``
  so one bad file cannot crash a batch of resumes.
* Everything works on in-memory bytes, so uploaded files are never written to disk.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}
MAX_FILE_BYTES = 10 * 1024 * 1024  # 10 MB per file
MIN_TEXT_CHARS = 50  # below this we treat the document as "no readable text"
MAX_PDF_PAGES = 40  # resumes are short; protects against huge files


class DocumentError(Exception):
    """Raised internally for any validation / extraction problem."""


@dataclass
class ParsedDocument:
    """Result of parsing one uploaded file."""

    filename: str
    text: str = ""
    num_pages: int | None = None
    error: str | None = None
    warnings: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return self.error is None and bool(self.text.strip())


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------
def validate_upload(filename: str, data: bytes) -> str:
    """Check extension, size and file signature. Returns the lower-case extension.

    Raises:
        DocumentError: with a user-friendly message.
    """
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise DocumentError(
            f"Unsupported file type '{ext or 'none'}'. Only PDF and DOCX are supported."
        )
    if not data:
        raise DocumentError("File is empty (0 bytes).")
    if len(data) > MAX_FILE_BYTES:
        mb = len(data) / (1024 * 1024)
        raise DocumentError(
            f"File is too large ({mb:.1f} MB). Limit is {MAX_FILE_BYTES // (1024 * 1024)} MB."
        )
    if ext == ".pdf" and b"%PDF-" not in data[:1024]:
        raise DocumentError("File is named .pdf but does not look like a valid PDF.")
    if ext == ".docx" and not (data[:2] == b"PK" and zipfile.is_zipfile(io.BytesIO(data))):
        raise DocumentError("File is named .docx but is not a valid Word (.docx) file.")
    return ext


# ---------------------------------------------------------------------------
# Extraction
# ---------------------------------------------------------------------------
def extract_pdf_text(data: bytes) -> tuple[str, int, list[str]]:
    """Extract text from PDF bytes using PyMuPDF. Returns (text, n_pages, warnings)."""
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:  # pragma: no cover - depends on environment
        raise DocumentError("PyMuPDF is not installed. Run: pip install PyMuPDF") from exc

    warnings: list[str] = []
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise DocumentError(f"PDF is corrupt or unreadable ({type(exc).__name__}).") from exc

    with doc:
        if doc.needs_pass:
            raise DocumentError("PDF is password-protected.")
        n_pages = doc.page_count
        if n_pages == 0:
            raise DocumentError("PDF has no pages.")
        pages_to_read = min(n_pages, MAX_PDF_PAGES)
        if n_pages > MAX_PDF_PAGES:
            warnings.append(f"Only the first {MAX_PDF_PAGES} of {n_pages} pages were read.")
        parts: list[str] = []
        try:
            for index in range(pages_to_read):
                parts.append(doc[index].get_text("text", sort=True))
        except Exception as exc:
            raise DocumentError(f"Failed while reading PDF pages ({type(exc).__name__}).") from exc
    return "\n".join(parts), n_pages, warnings


def extract_docx_text(data: bytes) -> tuple[str, list[str]]:
    """Extract text from DOCX bytes (paragraphs + table cells). Returns (text, warnings)."""
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - depends on environment
        raise DocumentError("python-docx is not installed. Run: pip install python-docx") from exc

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError(f"DOCX is corrupt or unreadable ({type(exc).__name__}).") from exc

    lines: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                value = cell.text.strip()
                if value and value not in cells:  # merged cells repeat their text
                    cells.append(value)
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines), []


def parse_document(filename: str, data: bytes) -> ParsedDocument:
    """Validate and parse one file. Never raises; errors are stored on the result."""
    result = ParsedDocument(filename=filename)
    try:
        ext = validate_upload(filename, data)
        if ext == ".pdf":
            text, n_pages, warnings = extract_pdf_text(data)
            result.num_pages = n_pages
        else:
            text, warnings = extract_docx_text(data)
        result.warnings.extend(warnings)

        stripped = text.strip()
        if len(stripped) < MIN_TEXT_CHARS:
            hint = (
                " It may be a scanned image PDF (OCR is not supported in this version)."
                if ext == ".pdf"
                else ""
            )
            raise DocumentError(f"No readable text found.{hint}")
        result.text = stripped
    except DocumentError as exc:
        result.error = str(exc)
    except Exception as exc:  # last-resort safety net
        result.error = f"Unexpected error while parsing ({type(exc).__name__}: {exc})."
    return result


def parse_path(path: str | Path) -> ParsedDocument:
    """Convenience wrapper for parsing a file on disk (used by scripts and tests)."""
    path = Path(path)
    try:
        data = path.read_bytes()
    except OSError as exc:
        return ParsedDocument(filename=path.name, error=f"Could not read file: {exc}")
    return parse_document(path.name, data)
